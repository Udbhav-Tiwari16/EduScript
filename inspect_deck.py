import sys
from pptx import Presentation

prs = Presentation('EduScript_Final_Project_Review_3.pptx')
with open('inspected_pptx_slides.txt', 'w', encoding='utf-8') as f:
    f.write(f'Total slides: {len(prs.slides)}\n')
    for idx, slide in enumerate(prs.slides):
        f.write(f'\n==============================\nSLIDE {idx+1}\n==============================\n')
        for s_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                f.write(f'  [Shape {s_idx} TextFrame] left={shape.left/914400:.2f}in, top={shape.top/914400:.2f}in, w={shape.width/914400:.2f}in, h={shape.height/914400:.2f}in\n')
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    txt = p.text.strip()
                    if txt:
                        runs_info = ', '.join([f'{r.text!r}(size={r.font.size.pt if r.font.size else None}, bold={r.font.bold})' for r in p.runs]) if p.runs else ''
                        f.write(f'    P{p_idx}: {txt}\n')
                        if runs_info:
                            f.write(f'       Runs: {runs_info}\n')
            elif shape.has_table:
                f.write(f'  [Shape {s_idx} Table] {len(shape.table.rows)} rows x {len(shape.table.columns)} cols\n')
                for r_idx, row in enumerate(shape.table.rows):
                    row_txt = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    f.write(f'    Row {r_idx}: {" | ".join(row_txt)}\n')
            elif shape.shape_type == 13: # picture
                f.write(f'  [Shape {s_idx} Picture] left={shape.left/914400:.2f}in, top={shape.top/914400:.2f}in, w={shape.width/914400:.2f}in, h={shape.height/914400:.2f}in\n')
            else:
                f.write(f'  [Shape {s_idx} ShapeType={shape.shape_type}] left={shape.left/914400:.2f}in, top={shape.top/914400:.2f}in\n')

print('Successfully dumped slide structure to inspected_pptx_slides.txt')
