"""
build_phase3_presentation.py
============================
Generates the official Phase 3 presentation:
"EduScript_Phase_3_Final_Implementation_and_Testing.pptx"

Preserves the exact visual identity, colors, typography, card structures,
and layout of the user's edited presentation, while using clean cropped
terminal evidence without window buttons, frames, or title bars.
"""

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


# -----------------------------------------------------------------------------
# Color Palette Constants (Exact match with user's edited template)
# -----------------------------------------------------------------------------
NAVY_DARK    = RGBColor(15, 23, 42)     # #0f172a
NAVY_LIGHT   = RGBColor(30, 41, 59)     # #1e293b
BLUE_PRIMARY = RGBColor(37, 99, 235)    # #2563eb
BLUE_LIGHT   = RGBColor(239, 246, 255)  # #eff6ff
GREEN_ACCENT = RGBColor(16, 185, 129)   # #10b981
GREEN_LIGHT  = RGBColor(236, 253, 245)  # #ecfdf5
PURPLE_ACC   = RGBColor(124, 58, 237)   # #7c3aed
TEXT_DARK    = RGBColor(15, 23, 42)     # #0f172a
TEXT_MUTED   = RGBColor(100, 116, 139)  # #64748b
TEXT_LIGHT   = RGBColor(248, 250, 252)  # #f8fafc
BG_LIGHT     = RGBColor(248, 250, 252)  # #f8fafc
CARD_BG      = RGBColor(255, 255, 255)  # #ffffff
CARD_BORDER  = RGBColor(226, 232, 240)  # #e2e8f0
CARD_BG_ALT  = RGBColor(241, 245, 249)  # #f1f5f9
RED_ACCENT   = RGBColor(239, 68, 68)    # #ef4444


def create_base_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_slide_header(slide, title_text: str, category_text: str = "EDUSCRIPT COMPILER DESIGN • PHASE 3: FINAL IMPLEMENTATION & TESTING", badge_text: str = "✓ PHASE 3 COMPLETE"):
    """Adds standard top bar with category tag, title, and status badge."""
    # Top accent bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.06))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = BLUE_PRIMARY
    top_bar.line.color.rgb = BLUE_PRIMARY

    # Category Tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.5), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = BLUE_PRIMARY

    # Phase 3 Badge on top right
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(0.5), Inches(2.333), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = GREEN_LIGHT
    badge.line.color.rgb = GREEN_ACCENT
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = badge_text
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = GREEN_ACCENT
    p_b.alignment = PP_ALIGN.CENTER

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.733), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY_DARK


def add_card(slide, left: float, top: float, width: float, height: float, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Adds a container card with smooth borders."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card


def add_bullet_point(tf, bold_prefix: str, text: str, font_size: int = 13, text_color=TEXT_DARK, pt_space: int = 6):
    p = tf.add_paragraph()
    p.space_after = Pt(pt_space)
    
    r1 = p.add_run()
    r1.text = bold_prefix + " "
    r1.font.bold = True
    r1.font.size = Pt(font_size)
    r1.font.color.rgb = text_color
    
    r2 = p.add_run()
    r2.text = text
    r2.font.bold = False
    r2.font.size = Pt(font_size)
    r2.font.color.rgb = text_color


# =============================================================================
# SLIDE BUILDERS (PHASE 3 SUBMISSION)
# =============================================================================

def build_slide_1_title(prs: Presentation):
    """Slide 1: Phase 3 Title Slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Dark Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_DARK
    bg.line.fill.background()

    # Top Accent Glow
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE_PRIMARY
    accent.line.fill.background()

    # Tag Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.4), Inches(5.8), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    badge.line.color.rgb = BLUE_PRIMARY
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "COMPILER DESIGN LAB • PHASE 3: FINAL IMPLEMENTATION & TESTING"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(191, 219, 254)
    p_b.alignment = PP_ALIGN.CENTER

    # Main Project Title
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(1.8))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "EduScript: A Customizable Lightweight Compiler and Interpreter"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    # Subtitle / Pipeline statement
    sbox = slide.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(11.333), Inches(0.8))
    tf_s = sbox.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = "End-to-End Compiler Pipeline: Lexical Analysis → Recursive Descent Parser → AST → Scoped Symbol Table → Semantic Analyzer → Three-Address Code (TAC) → Virtual Machine Interpreter"
    p_s.font.size = Pt(13.5)
    p_s.font.color.rgb = RGBColor(148, 163, 184)

    # Info Cards at Bottom
    cards = [
        ("STUDENT NAME", "Udbhav Tiwari", BLUE_PRIMARY),
        ("REGISTER NUMBER", "24BCT0328", PURPLE_ACC),
        ("COURSE & SUBMISSION", "CD Lab • Phase 3 (Final)", GREEN_ACCENT),
        ("PROJECT TYPE", "Individual Project", RGBColor(234, 88, 12))
    ]

    for idx, (label, val, col) in enumerate(cards):
        cx = 1.0 + idx * 2.9
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(5.1), Inches(2.65), Inches(1.6))
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY_LIGHT
        c.line.color.rgb = RGBColor(51, 65, 85)
        
        # Indicator bar
        ibar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(5.1), Inches(2.65), Inches(0.06))
        ibar.fill.solid()
        ibar.fill.fore_color.rgb = col
        ibar.line.fill.background()

        tf_c = c.text_frame
        tf_c.word_wrap = True
        p_l = tf_c.paragraphs[0]
        p_l.text = label
        p_l.font.size = Pt(10)
        p_l.font.bold = True
        p_l.font.color.rgb = RGBColor(148, 163, 184)
        p_l.space_before = Pt(8)

        p_v = tf_c.add_paragraph()
        p_v.text = val
        p_v.font.size = Pt(14)
        p_v.font.bold = True
        p_v.font.color.rgb = TEXT_LIGHT
        p_v.space_before = Pt(4)


def build_slide_2_overview(prs: Presentation):
    """Slide 2: Phase 3 Objective & Project Completion."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 3 Objectives & Final System Overview")

    # Left Card: Core Pipeline Stages
    add_card(slide, 0.8, 1.6, 5.6, 5.3)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(5.2), Inches(5.0))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Complete 5-Stage Compiler Pipeline"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = BLUE_PRIMARY
    p0.space_after = Pt(10)

    add_bullet_point(tf, "1. Lexical Analysis (lexer.py):", "Tokenizes raw source with exact line/col tracking and normalizes custom keywords (agar, jabtak, purnaank, etc.) via JSON mappings.")
    add_bullet_point(tf, "2. Syntax Analysis (parser.py):", "Recursive descent parsing with full precedence climbing; constructs structured Abstract Syntax Trees with panic-mode recovery.")
    add_bullet_point(tf, "3. Semantic Analysis (semantic_analyzer.py):", "Enforces static type checking, scoped symbol table lookups, variable shadowing, and diagnostic validation.")
    add_bullet_point(tf, "4. Intermediate Code (tac_generator.py):", "Emits linearized Three-Address Code (TAC) with temporary registers (t0, t1...) and jump labels (L0, L1...).")
    add_bullet_point(tf, "5. Execution Engine (interpreter.py):", "Custom TAC Virtual Machine instruction dispatcher, plus auxiliary AST Tree-Walk execution mode.")

    # Right Card: Phase 3 Deliverables & Highlights
    add_card(slide, 6.8, 1.6, 5.733, 5.3)
    tbox_r = slide.shapes.add_textbox(Inches(7.0), Inches(1.75), Inches(5.333), Inches(5.0))
    tf_r = tbox_r.text_frame
    tf_r.word_wrap = True

    pr0 = tf_r.paragraphs[0]
    pr0.text = "Phase 3 Implementation & Testing Highlights"
    pr0.font.size = Pt(16)
    pr0.font.bold = True
    pr0.font.color.rgb = GREEN_ACCENT
    pr0.space_after = Pt(10)

    add_bullet_point(tf_r, "• 100% Implemented & Tested:", "Zero mockups or placeholders; fully working end-to-end interpreter executing real programs.")
    add_bullet_point(tf_r, "• Multilingual/Custom Keywords:", "Supports native Hindi keywords (purnaank, agar, warna, jabtak, dikhao, wapas) mapped to canonical constructs.")
    add_bullet_point(tf_r, "• Comprehensive Test Suite:", "12 automated test cases covering valid programs, Hindi keywords, arithmetic, for/while loops, and error diagnostics.")
    add_bullet_point(tf_r, "• Academic Transparency:", "CLI flags (--tokens, --ast, --symbols, --tac, --all) allow live inspection of every single compiler transformation stage.")
    add_bullet_point(tf_r, "• Zero External Dependencies:", "Built natively with Python standard library for maximum portability and educational clarity.")


def build_slide_3_architecture(prs: Presentation):
    """Slide 3: System Architecture & Inter-Module Data Flow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "System Architecture & Inter-Module Data Flow")

    steps = [
        ("STAGE 1: SCANNER", "lexer.py", "Input: .custom / .edu source\nOutput: List[Token]\n• JSON keyword mapping\n• Exact Line:Col tracking\n• Non-fatal lexical errors", BLUE_PRIMARY),
        ("STAGE 2: PARSER", "parser.py & ast_nodes.py", "Input: List[Token]\nOutput: Program AST Node\n• Recursive descent parser\n• Operator precedence\n• Panic-mode recovery", PURPLE_ACC),
        ("STAGE 3: SEMANTIC", "symbol_table.py & analyzer", "Input: Program AST\nOutput: Validated AST + SymTab\n• Lexical scope hierarchy\n• Static type compatibility\n• Scope & symbol checking", RGBColor(234, 88, 12)),
        ("STAGE 4 & 5: IR & VM", "tac_generator & interpreter", "Input: Validated AST\nOutput: Execution & Output\n• 3-Address Code generation\n• Virtual Machine execution\n• Runtime safety & output", GREEN_ACCENT)
    ]

    for idx, (title, mod, desc, col) in enumerate(steps):
        left = 0.8 + idx * 2.98
        add_card(slide, left, 1.6, 2.8, 5.3)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(1.6), Inches(2.8), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tbox = slide.shapes.add_textbox(Inches(left + 0.15), Inches(1.75), Inches(2.5), Inches(5.0))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = mod
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(10)

        for line in desc.splitlines():
            p_desc = tf.add_paragraph()
            p_desc.text = line
            p_desc.font.size = Pt(11)
            p_desc.font.color.rgb = TEXT_DARK
            p_desc.space_after = Pt(4)


def build_slide_4_lexer(prs: Presentation, screenshots_dir: Path):
    """Slide 4: Phase 1 Lexical Analysis & Custom Keywords."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 1: Lexical Analysis & Custom Keyword Normalization")

    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Lexer Design & Keyword Normalization"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = BLUE_PRIMARY
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Data-Driven Lexer (lexer.py):", "Loads token specifications from default_tokens.json and keyword aliases from keyword_mapping.json.")
    add_bullet_point(tf, "• Transparent Normalization:", "Custom words (purnaank, jabtak, agar) are converted into canonical tokens (int, while, if) while preserving original lexemes.")
    add_bullet_point(tf, "• Longest-Match Operators:", "Multi-character operators (==, !=, <=, >=, +=, -=) sorted by length to match before single-character operators (=, <).")
    add_bullet_point(tf, "• Position Tracking:", "Tracks exact 1-based line and column coordinates for every token, enabling pinpoint compiler error messages.")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "03_phase1_tokens.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.41), width=Inches(6.15))


def build_slide_5_parser_ast(prs: Presentation, screenshots_dir: Path):
    """Slide 5: Phase 2 Syntax Analysis & AST."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 2: Syntax Analysis & Abstract Syntax Tree (AST)")

    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Recursive Descent Parser (parser.py)"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = PURPLE_ACC
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Precedence Climbing:", "Mathematical operator precedence enforced directly through grammar function hierarchy without shift-reduce ambiguities.")
    add_bullet_point(tf, "• Modular AST Nodes (ast_nodes.py):", "Strongly-typed nodes for Program, VarDecl, IfStmt, WhileStmt, ForStmt, BinaryExpr, AssignExpr, etc.")
    add_bullet_point(tf, "• Visitor Pattern Support:", "AST nodes accept visitors (accept(self, visitor)) enabling clean decoupled passes for Semantics, TAC IR, and Pretty Printing.")
    add_bullet_point(tf, "• Panic-Mode Recovery:", "Synchronizes at statement boundaries (semicolons, closing braces) to catch subsequent syntax errors without crashing.")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "04_phase2_ast.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.41), width=Inches(6.15))


def build_slide_6_symbol_table_semantics(prs: Presentation, screenshots_dir: Path):
    """Slide 6: Phase 3 Scoped Symbol Table & Semantic Analysis."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 3: Scoped Symbol Table & Semantic Analysis")

    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Semantic Validation & Scope Engine"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(234, 88, 12)
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Scoped Symbol Table (symbol_table.py):", "Nested Scope tree maintaining identifier name, data type (int, float, string, bool), scope level, and initialization status.")
    add_bullet_point(tf, "• Variable Shadowing:", "Inner block scopes can shadow outer symbols; symbol lookups climb the parent scope chain upwards to global scope.")
    add_bullet_point(tf, "• Type Checking (semantic_analyzer.py):", "Enforces static type compatibility; allows widening int -> float; validates relational & logical operands.")
    add_bullet_point(tf, "• Semantic Error Trapping:", "Flags duplicate declarations in identical scope, undeclared identifier usage, and uninitialized variable reads.")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "05_phase3_symbol_table.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.85), width=Inches(6.15))


def build_slide_7_tac_ir(prs: Presentation, screenshots_dir: Path):
    """Slide 7: Phase 4 Intermediate Representation (TAC)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 4: Intermediate Code Generation (Three-Address Code)")

    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Three-Address Code (tac_generator.py)"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = BLUE_PRIMARY
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Linearized Quadruple IR:", "Lowers high-level nested AST nodes into elementary 3-address instructions with at most 1 operator per instruction.")
    add_bullet_point(tf, "• Register & Label Generation:", "Emits temporary registers (t0, t1, t2...) for intermediate sub-expressions and symbolic jump labels (L_while_start_0, L_else_2...).")
    add_bullet_point(tf, "• Complete Instruction Set:", "Supports ASSIGN, ADD, SUB, MUL, DIV, MOD, EQ, NEQ, LT, LE, GT, GE, AND, OR, NOT, GOTO, IF_FALSE, IF_TRUE, PRINT, RETURN.")
    add_bullet_point(tf, "• Clear Machine Abstraction:", "Isolates front-end language syntax from the underlying execution runtime.")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "06_phase4_tac_ir.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.41), width=Inches(6.15))


def build_slide_8_interpreter(prs: Presentation, screenshots_dir: Path):
    """Slide 8: Phase 5 Execution Engine & Virtual Machine."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 5: Execution Engine & TAC Virtual Machine")

    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Execution Engine (interpreter.py)"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = GREEN_ACCENT
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• TAC Virtual Machine (Primary):", "Dispatches TAC instructions step-by-step with program counter management, environment register state, and O(1) label index lookups.")
    add_bullet_point(tf, "• AST Tree-Walk Interpreter (Auxiliary):", "Direct AST traversal execution engine available via --ast-run flag for comparative demonstration.")
    add_bullet_point(tf, "• Runtime Error Trapping:", "Gracefully catches and reports division/modulo by zero and guards against infinite execution loops via max-step thresholds.")
    add_bullet_point(tf, "• Program Exit & Output:", "Captures stdout cleanly in execution buffers and propagates exit return codes (0 on success).")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "07_hindi_factorial_run.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(2.1), width=Inches(6.15))


def build_slide_9_end_to_end(prs: Presentation, screenshots_dir: Path):
    """Slide 9: End-to-End Pipeline Demonstration."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Live End-to-End Execution Evidence (All 5 Stages)")

    add_card(slide, 0.8, 1.24, 4.8, 5.66)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(4.4), Inches(5.2))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Full Pipeline Execution: example.custom"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = BLUE_PRIMARY
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Source Code Tested:", "Custom Hindi keywords script executing while loops (jabtak), float arithmetic (dashamlav), and conditional branching (agar/warna).")
    add_bullet_point(tf, "• Executed Command:", "python main.py example.custom --all")
    add_bullet_point(tf, "• Complete Pipeline Verification:", "Simultaneously outputs Token Stream (56 tokens) → AST Tree → Scoped Symbol Table → TAC (20 instructions) → Program Output ('High score achieved!').")
    add_bullet_point(tf, "• Verification Status:", "Zero lexical, syntax, or semantic errors; exited with return code 0.")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "02_full_pipeline_all.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(5.95), Inches(1.24), width=Inches(6.55))


def build_slide_10_error_diagnostics(prs: Presentation, screenshots_dir: Path):
    """Slide 10: Multi-Phase Error Diagnostics."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Robust Multi-Phase Error Diagnostic Architecture")

    errors_info = [
        ("1. LEXICAL ERROR (Phase 1)", "Unsupported character '$ @ ~' caught with line/col location.", "12_error_lexical.png", 0.8, 1.6),
        ("2. SYNTAX ERROR (Phase 2)", "Malformed expression 'if (a > 5 {' caught by recursive parser.", "11_error_syntax.png", 6.8, 1.6),
        ("3. SEMANTIC ERROR (Phase 3)", "Type mismatch assigning string to integer variable 'count'.", "09_error_type_mismatch.png", 0.8, 4.3),
        ("4. RUNTIME ERROR (Phase 5)", "Division by zero 'a / b' safely trapped at VM execution time.", "13_error_division_by_zero.png", 6.8, 4.3)
    ]

    for title, desc, img_name, left, top in errors_info:
        add_card(slide, left, top, 5.733, 2.6)
        tbox = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.1), Inches(5.4), Inches(0.55))
        tf = tbox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = RED_ACCENT

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

        img_path = screenshots_dir / img_name
        if img_path.exists():
            slide.shapes.add_picture(str(img_path), Inches(left + 0.15), Inches(top + 0.72), width=Inches(5.433))


def build_slide_11_test_suite(prs: Presentation, screenshots_dir: Path):
    """Slide 11: Automated Test Suite & Results."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 3 Testing Strategy & 100% Verification Proof")

    add_card(slide, 0.8, 1.41, 5.5, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.1), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Verified 12-Case Test Matrix"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = GREEN_ACCENT
    p0.space_after = Pt(8)

    test_rows = [
        ("1. Custom Hindi (example.custom)", "Loops, floats, if/else", "PASSED"),
        ("2. Hindi Factorial", "Loop accumulator, compare", "PASSED"),
        ("3. Arithmetic & Precedence", "Ops (+,-,*,/,%), comp. assign", "PASSED"),
        ("4. Control Flow & Logic", "While, modulo, &&, ||", "PASSED"),
        ("5. For Loop", "Init, cond, update loop", "PASSED"),
        ("6. String Operations", "Concatenation, equality", "PASSED"),
        ("7. Semantic: Type Mismatch", "Exit code 3 verified", "PASSED"),
        ("8. Semantic: Undeclared Var", "Exit code 3 verified", "PASSED"),
        ("9. Semantic: Duplicate Decl", "Exit code 3 verified", "PASSED"),
        ("10. Syntax: Malformed Expr", "Exit code 2 verified", "PASSED"),
        ("11. Lexical: Bad Characters", "Exit code 1 verified", "PASSED"),
        ("12. Runtime: Zero Division", "Exit code 4 verified", "PASSED")
    ]

    for name, desc, status in test_rows:
        p = tf.add_paragraph()
        p.space_after = Pt(2)
        r1 = p.add_run()
        r1.text = f"✓ {name}: "
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_DARK
        
        r2 = p.add_run()
        r2.text = f"{desc} "
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

        r3 = p.add_run()
        r3.text = f"[{status}]"
        r3.font.bold = True
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = GREEN_ACCENT

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "01_test_suite_runner.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.5), Inches(1.41), width=Inches(6.033))


def build_slide_12_planned_vs_completed(prs: Presentation):
    """Slide 12: Planned System vs Final Delivered System."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Compliance Matrix: Planned Proposal vs Final Implementation")

    rows = [
        ("Component", "Initial Proposal Target", "Phase 3 Final Delivery", "Status"),
        ("Lexical Analyzer", "Data-driven tokenization with custom keywords", "Dynamic JSON-driven lexer with line/col tracking & Hindi mappings", "100% COMPLETE"),
        ("Parser", "Recursive descent parser for imperative constructs", "Full recursive descent with precedence climbing & error recovery", "100% COMPLETE"),
        ("AST", "Modular AST node hierarchy with visitor pattern", "Typed AST nodes with visitor interface & visual ASCII tree printer", "100% COMPLETE"),
        ("Symbol Table", "Scoped variable table with type records", "Hierarchical Scope trees, variable shadowing & initialization status", "100% COMPLETE"),
        ("Semantic Analyzer", "Type checking & undeclared identifier detection", "Static type inference, widening int->float, duplicate decl checks", "100% COMPLETE"),
        ("Intermediate Code", "Three-Address Code (TAC) generation", "Linearized TAC quadruples with temporaries (t0..) and labels (L0..)", "100% COMPLETE"),
        ("Execution Engine", "Interpreter executing EduScript code", "TAC Virtual Machine + auxiliary AST Tree-Walk Interpreter", "100% COMPLETE"),
        ("Testing Suite", "Automated validation of language features", "12 automated test cases with 100.0% pass rate & CI-ready runner", "100% COMPLETE")
    ]

    left = Inches(0.8)
    top = Inches(1.6)
    width = Inches(11.733)
    height = Inches(5.3)

    table_shape = slide.shapes.add_table(len(rows), 4, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.6)
    table.columns[2].width = Inches(4.3)
    table.columns[3].width = Inches(1.633)

    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row):
            tc = table.cell(r_idx, c_idx)
            tc.text = cell
            p = tc.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c_idx == 3 else PP_ALIGN.LEFT
            
            if r_idx == 0:
                p.font.bold = True
                p.font.size = Pt(11)
                p.font.color.rgb = TEXT_LIGHT
                tc.fill.solid()
                tc.fill.fore_color.rgb = NAVY_DARK
            else:
                p.font.size = Pt(10)
                if c_idx == 3:
                    p.font.bold = True
                    p.font.color.rgb = GREEN_ACCENT
                elif c_idx == 0:
                    p.font.bold = True
                    p.font.color.rgb = NAVY_DARK
                else:
                    p.font.color.rgb = TEXT_DARK
                tc.fill.solid()
                tc.fill.fore_color.rgb = CARD_BG_ALT if r_idx % 2 == 1 else CARD_BG


def build_slide_13_technical_innovation(prs: Presentation):
    """Slide 13: Technical Innovation & Engineering Contributions."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Technical Innovations & Engineering Highlights")

    innovations = [
        ("1. Dynamic JSON Keyword Customization", "Enables localized / multilingual programming (e.g. Hindi keywords agar, jabtak, purnaank) without altering parser grammar rules. The Lexer normalizes custom aliases to canonical tags transparently.", BLUE_PRIMARY),
        ("2. Total Academic Transparency", "Unlike black-box compilers, every intermediate stage (Tokens, AST, Symbol Table, TAC IR) is fully inspectable via dedicated CLI flags (--tokens, --ast, --symbols, --tac, --all).", PURPLE_ACC),
        ("3. Dual-Execution Architecture", "Features both a Three-Address Code Virtual Machine (executing linearized IR) and an AST Tree-Walk Interpreter, allowing direct comparison of IR vs AST execution in compiler vivas.", GREEN_ACCENT),
        ("4. Pinpoint Diagnostic Architecture", "Multi-phase error trapping with exact line and column coordinates, detailed error classification (Lexical, Syntax, TypeError, UndeclaredVariable, RuntimeError), and panic-mode recovery.", RGBColor(234, 88, 12))
    ]

    for idx, (title, desc, col) in enumerate(innovations):
        row = idx // 2
        col_idx = idx % 2
        left = 0.8 + col_idx * 5.95
        top = 1.6 + row * 2.7

        add_card(slide, left, top, 5.75, 2.5)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(5.75), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tbox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(5.35), Inches(2.2))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK


def build_slide_14_git_deliverables(prs: Presentation, screenshots_dir: Path):
    """Slide 14: Phase 3 Deliverables & GitHub Repository."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Phase 3 Deliverables & GitHub Repository")

    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Repository & Deliverables Summary"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = NAVY_DARK
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• GitHub Repository:", "https://github.com/Udbhav-Tiwari16/EduScript")
    add_bullet_point(tf, "• Active Branch:", "main (All changes committed & pushed)")
    add_bullet_point(tf, "• Phase 3 Deliverables Complete:", "✓ Complete Working Project\n✓ Complete Source Code\n✓ Test Cases & Verified Results\n✓ Screenshots / Output Evidence\n✓ Final Presentation (.pptx)\n✓ User Guide & Documentation")
    add_bullet_point(tf, "• Individual Project:", "Sole author and developer: Udbhav Tiwari (Reg No: 24BCT0328).")

    # Clean Cropped Terminal Picture
    img_path = screenshots_dir / "14_git_status_log.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(2.6), width=Inches(6.15))


def build_slide_15_conclusion(prs: Presentation):
    """Slide 15: Conclusion & Future Scope."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Conclusion & Future Enhancements")

    # Left: Conclusion
    add_card(slide, 0.8, 1.6, 5.6, 5.3)
    tbox_l = slide.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(5.2), Inches(5.0))
    tf_l = tbox_l.text_frame
    tf_l.word_wrap = True

    pl0 = tf_l.paragraphs[0]
    pl0.text = "Phase 3 Conclusion"
    pl0.font.size = Pt(16)
    pl0.font.bold = True
    pl0.font.color.rgb = GREEN_ACCENT
    pl0.space_after = Pt(8)

    add_bullet_point(tf_l, "• Comprehensive Implementation:", "EduScript fulfills all theoretical and practical requirements of a complete compiler/interpreter pipeline.")
    add_bullet_point(tf_l, "• Educational & Modular:", "Every stage is decoupled, inspectable, and built in pure Python standard library for maximum readability.")
    add_bullet_point(tf_l, "• Fully Verified:", "100% test pass rate across 12 automated test cases covering valid programs, custom keywords, and error handling.")
    add_bullet_point(tf_l, "• Ready for Evaluation:", "Complete source code committed and pushed to GitHub repository.")

    # Right: Future Scope
    add_card(slide, 6.8, 1.6, 5.733, 5.3)
    tbox_r = slide.shapes.add_textbox(Inches(7.0), Inches(1.75), Inches(5.333), Inches(5.0))
    tf_r = tbox_r.text_frame
    tf_r.word_wrap = True

    pr0 = tf_r.paragraphs[0]
    pr0.text = "Future Scope & Enhancements"
    pr0.font.size = Pt(16)
    pr0.font.bold = True
    pr0.font.color.rgb = BLUE_PRIMARY
    pr0.space_after = Pt(8)

    add_bullet_point(tf_r, "• TAC Optimization Passes:", "Dead code elimination, constant folding, constant propagation, and common subexpression elimination (CSE).")
    add_bullet_point(tf_r, "• Bytecode / Assembly Target:", "Compiling TAC to Python bytecode (.pyc) or LLVM IR for native machine execution.")
    add_bullet_point(tf_r, "• User-Defined Functions & Arrays:", "Expanding runtime stack frames to support recursive function calls and array data structures.")
    add_bullet_point(tf_r, "• Interactive Web Playground:", "WebAssembly frontend allowing students to visualize AST trees and TAC steps interactively in a browser.")


def main():
    base_dir = Path(__file__).parent.resolve()
    screenshots_dir = base_dir / "evidence_screenshots"
    output_ppt_path = base_dir / "EduScript_Phase_3_Final_Implementation_and_Testing.pptx"
    output_root_ppt_path = base_dir.parent / "EduScript_Phase_3_Final_Implementation_and_Testing.pptx"

    prs = create_base_presentation()

    print("Building Phase 3 presentation slides...")
    build_slide_1_title(prs)
    build_slide_2_overview(prs)
    build_slide_3_architecture(prs)
    build_slide_4_lexer(prs, screenshots_dir)
    build_slide_5_parser_ast(prs, screenshots_dir)
    build_slide_6_symbol_table_semantics(prs, screenshots_dir)
    build_slide_7_tac_ir(prs, screenshots_dir)
    build_slide_8_interpreter(prs, screenshots_dir)
    build_slide_9_end_to_end(prs, screenshots_dir)
    build_slide_10_error_diagnostics(prs, screenshots_dir)
    build_slide_11_test_suite(prs, screenshots_dir)
    build_slide_12_planned_vs_completed(prs)
    build_slide_13_technical_innovation(prs)
    build_slide_14_git_deliverables(prs, screenshots_dir)
    build_slide_15_conclusion(prs)

    prs.save(str(output_ppt_path))
    prs.save(str(output_root_ppt_path))
    print(f"\nSuccessfully generated Phase 3 presentation:\n- {output_ppt_path}\n- {output_root_ppt_path}")


if __name__ == "__main__":
    main()
