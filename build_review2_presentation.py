"""
build_review2_presentation.py
=============================
Generates the official Review 2 presentation:
"EduScript_Review_2_Implementation_Progress.pptx"

Strictly scoped to Review 2: Implementation Progress Review
(8 Official Review 2 Criteria: Implementation Progress, Functional Correctness,
Compiler Concepts, Code Quality, Testing, Problem Solving, Innovation, Viva).
10 focused slides preserving the exact visual identity of the edited template.
"""

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


# -----------------------------------------------------------------------------
# Color Palette Constants (Exact match with user's edited template)
# -----------------------------------------------------------------------------
NAVY_DARK    = RGBColor(15, 23, 42)     # #0f172a
NAVY_LIGHT   = RGBColor(30, 41, 59)     # #1e293b
BLUE_PRIMARY = RGBColor(37, 99, 235)    # #2563eb
BLUE_LIGHT   = RGBColor(239, 246, 255)  # #eff6ff
GREEN_ACCENT = RGBColor(16, 185, 129)   # #10b981
GREEN_LIGHT  = RGBColor(236, 253, 245)  # #ecfdf5
PURPLE_ACC   = RGBColor(124, 58, 237)   # #7c3aed
TEXT_DARK    = RGBColor(15, 23, 42)     # #0f172a
TEXT_MUTED   = RGBColor(100, 116, 139)  # #64748b
TEXT_LIGHT   = RGBColor(248, 250, 252)  # #f8fafc
BG_LIGHT     = RGBColor(248, 250, 252)  # #f8fafc
CARD_BG      = RGBColor(255, 255, 255)  # #ffffff
CARD_BORDER  = RGBColor(226, 232, 240)  # #e2e8f0
CARD_BG_ALT  = RGBColor(241, 245, 249)  # #f1f5f9
ORANGE_ACC   = RGBColor(234, 88, 12)    # #ea580c


def create_base_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_slide_header(slide, title_text: str, category_text: str = "EDUSCRIPT COMPILER DESIGN • REVIEW 2: IMPLEMENTATION PROGRESS", badge_text: str = "REVIEW 2 PROGRESS"):
    """Adds standard top bar with category tag, title, and status badge."""
    # Top accent bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.06))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = BLUE_PRIMARY
    top_bar.line.color.rgb = BLUE_PRIMARY

    # Category Tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.5), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = BLUE_PRIMARY

    # Review 2 Badge on top right
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(0.5), Inches(2.333), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_PRIMARY
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = badge_text
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = BLUE_PRIMARY
    p_b.alignment = PP_ALIGN.CENTER

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.733), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY_DARK


def add_card(slide, left: float, top: float, width: float, height: float, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Adds a container card with smooth borders."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card


def add_bullet_point(tf, bold_prefix: str, text: str, font_size: int = 13, text_color=TEXT_DARK, pt_space: int = 6):
    p = tf.add_paragraph()
    p.space_after = Pt(pt_space)
    
    r1 = p.add_run()
    r1.text = bold_prefix + " "
    r1.font.bold = True
    r1.font.size = Pt(font_size)
    r1.font.color.rgb = text_color
    
    r2 = p.add_run()
    r2.text = text
    r2.font.bold = False
    r2.font.size = Pt(font_size)
    r2.font.color.rgb = text_color


# =============================================================================
# SLIDE BUILDERS (REVIEW 2: 10 FOCUSED SLIDES)
# =============================================================================

def build_slide_1_title(prs: Presentation):
    """Slide 1: Title Slide (Review 2)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Dark Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_DARK
    bg.line.fill.background()

    # Top Accent Line
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE_PRIMARY
    accent.line.fill.background()

    # Review 2 Tag Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.4), Inches(5.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    badge.line.color.rgb = BLUE_PRIMARY
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "COMPILER DESIGN LABORATORY • REVIEW 2"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(191, 219, 254)
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(1.8))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "EduScript: Implementation Progress Review"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    # Subtitle
    sbox = slide.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(11.333), Inches(0.8))
    tf_s = sbox.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = "Progress Report: Modular Architecture, Lexical & Syntax Pipeline, Scoped Symbol Management, and Working Components"
    p_s.font.size = Pt(13.5)
    p_s.font.color.rgb = RGBColor(148, 163, 184)

    # 4 Bottom Cards
    cards = [
        ("STUDENT NAME", "Udbhav Tiwari", BLUE_PRIMARY),
        ("REGISTER NUMBER", "24BCT0328", PURPLE_ACC),
        ("COURSE & REVIEW", "CD Lab • Review 2", GREEN_ACCENT),
        ("PROJECT TYPE", "Individual Project", ORANGE_ACC)
    ]

    for idx, (label, val, col) in enumerate(cards):
        cx = 1.0 + idx * 2.9
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(5.1), Inches(2.65), Inches(1.6))
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY_LIGHT
        c.line.color.rgb = RGBColor(51, 65, 85)
        
        ibar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(5.1), Inches(2.65), Inches(0.06))
        ibar.fill.solid()
        ibar.fill.fore_color.rgb = col
        ibar.line.fill.background()

        tf_c = c.text_frame
        tf_c.word_wrap = True
        p_l = tf_c.paragraphs[0]
        p_l.text = label
        p_l.font.size = Pt(10)
        p_l.font.bold = True
        p_l.font.color.rgb = RGBColor(148, 163, 184)
        p_l.space_before = Pt(8)

        p_v = tf_c.add_paragraph()
        p_v.text = val
        p_v.font.size = Pt(14)
        p_v.font.bold = True
        p_v.font.color.rgb = TEXT_LIGHT
        p_v.space_before = Pt(4)


def build_slide_2_overview(prs: Presentation):
    """Slide 2: Project Overview & Planned Work."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Project Overview & Planned System Architecture")

    # Left Card: Problem & Objectives
    add_card(slide, 0.8, 1.6, 5.6, 5.3)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(5.2), Inches(5.0))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Problem Statement & Objectives"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = BLUE_PRIMARY
    p0.space_after = Pt(10)

    add_bullet_point(tf, "• Purpose of EduScript:", "To build a transparent, highly educational compiler/interpreter for an imperative language with customizable syntax.")
    add_bullet_point(tf, "• Educational Gap Addressed:", "Standard production compilers act as black boxes; EduScript enables students to inspect every transformation stage in real time.")
    add_bullet_point(tf, "• Key Project Goals:", "1. Data-driven lexical scanning with custom multilingual keywords.\n2. Recursive descent syntax analysis with structured AST nodes.\n3. Scoped symbol tables with static semantic validation.\n4. Intermediate code generation & execution engine.")

    # Right Card: Modular Architecture
    add_card(slide, 6.8, 1.6, 5.733, 5.3)
    tbox_r = slide.shapes.add_textbox(Inches(7.0), Inches(1.75), Inches(5.333), Inches(5.0))
    tf_r = tbox_r.text_frame
    tf_r.word_wrap = True

    pr0 = tf_r.paragraphs[0]
    pr0.text = "Modular Pipeline Architecture"
    pr0.font.size = Pt(16)
    pr0.font.bold = True
    pr0.font.color.rgb = GREEN_ACCENT
    pr0.space_after = Pt(10)

    add_bullet_point(tf_r, "• Modular Staged Design:", "Each compiler phase is implemented as an independent, decoupled module communicating through explicit data structures.")
    add_bullet_point(tf_r, "• Lexical Analyzer (lexer.py):", "Converts raw characters into normalized Token stream using external JSON keyword mapping.")
    add_bullet_point(tf_r, "• Parser & AST (parser.py, ast_nodes.py):", "Validates grammatical structure and generates a strongly-typed Abstract Syntax Tree.")
    add_bullet_point(tf_r, "• Symbol Table & Semantic Checks:", "Tracks hierarchical lexical scopes, identifier declarations, and type compatibility.")
    add_bullet_point(tf_r, "• Backend Pipeline:", "Translates validated AST into Three-Address Code for virtual machine execution.")


def build_slide_3_progress(prs: Presentation):
    """Slide 3: Implementation Progress (Review 2 Scope)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Implementation Progress & Current Development Status")

    # Progress Table
    rows = [
        ("Compiler Module / Component", "Planned Functionality", "Implementation Status", "Verification Status"),
        ("Lexical Analyzer (lexer.py)", "Data-driven tokenization, keyword mapping, line/col tracking", "Implemented & Operational", "Verified with test suite"),
        ("Configuration System (JSON)", "External keyword mappings and token definitions", "Implemented & Configurable", "Tested with Hindi keywords"),
        ("Parser (parser.py)", "Recursive descent parsing with precedence climbing", "Implemented & Operational", "Verified with AST tests"),
        ("AST Hierarchy (ast_nodes.py)", "Typed AST node classes, Visitor pattern, ASCII printer", "Implemented & Modular", "Verified across statements"),
        ("Symbol Table (symbol_table.py)", "Hierarchical scope tree, variable shadowing & state", "Implemented & Scoped", "Tested across block scopes"),
        ("Semantic Analyzer", "Static type checking, undeclared & duplicate checks", "Implemented & Validated", "Diagnostic errors verified"),
        ("Intermediate Code & Execution", "Three-Address Code generation & Virtual Machine", "Implemented & Integrated", "End-to-end execution active")
    ]

    left = Inches(0.8)
    top = Inches(1.6)
    width = Inches(11.733)
    height = Inches(5.3)

    table_shape = slide.shapes.add_table(len(rows), 4, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(4.3)
    table.columns[2].width = Inches(2.5)
    table.columns[3].width = Inches(2.333)

    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row):
            tc = table.cell(r_idx, c_idx)
            tc.text = cell
            p = tc.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c_idx >= 2 else PP_ALIGN.LEFT
            
            if r_idx == 0:
                p.font.bold = True
                p.font.size = Pt(11)
                p.font.color.rgb = TEXT_LIGHT
                tc.fill.solid()
                tc.fill.fore_color.rgb = NAVY_DARK
            else:
                p.font.size = Pt(10)
                if c_idx == 2:
                    p.font.bold = True
                    p.font.color.rgb = BLUE_PRIMARY
                elif c_idx == 3:
                    p.font.bold = True
                    p.font.color.rgb = GREEN_ACCENT
                elif c_idx == 0:
                    p.font.bold = True
                    p.font.color.rgb = NAVY_DARK
                else:
                    p.font.color.rgb = TEXT_DARK
                tc.fill.solid()
                tc.fill.fore_color.rgb = CARD_BG_ALT if r_idx % 2 == 1 else CARD_BG


def build_slide_4_working_module(prs: Presentation, screenshots_dir: Path):
    """Slide 4: Working Module Demonstration (Lexical Analysis & Keyword Mapping)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Working Module Demonstration: Lexical Analyzer")

    # Left: Explanation
    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Lexer Execution & Keyword Normalization"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = BLUE_PRIMARY
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Working Command:", "python main.py example.custom --tokens")
    add_bullet_point(tf, "• Token Recognition:", "Scans raw source text into structured atomic Token objects with type, lexeme, and 1-based line:column coordinates.")
    add_bullet_point(tf, "• Custom Keyword Normalization:", "Custom Hindi keywords (purnaank, jabtak, agar, warna, dikhao, wapas) are transparently mapped to canonical equivalents (int, while, if, else, print, return).")
    add_bullet_point(tf, "• Token Stream Breakdown:", "Outputs clean tabular breakdown categorizing identifiers, integers, floats, operators, and delimiters.")

    # Right: Clean Cropped Terminal Evidence
    img_path = screenshots_dir / "03_phase1_tokens.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.41), width=Inches(6.15))


def build_slide_5_compiler_concepts(prs: Presentation, screenshots_dir: Path):
    """Slide 5: Compiler Concepts and Technical Implementation."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Compiler Concepts & Technical Implementation")

    # Left: Concepts applied
    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Application of Core Compiler Theory"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = PURPLE_ACC
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Longest-Match Tokenization:", "Operators sorted by length to guarantee multi-character operators (==, !=, <=, >=) match before single-character tokens (=, <).")
    add_bullet_point(tf, "• Recursive Descent Parsing:", "Deterministic top-down grammar parsing with precedence climbing ensuring correct arithmetic & relational operator binding.")
    add_bullet_point(tf, "• Structured AST Construction:", "Abstract Syntax Tree constructed with dedicated node classes, preserving hierarchy for subsequent analysis passes.")
    add_bullet_point(tf, "• Panic-Mode Error Recovery:", "Synchronizes at statement delimiters (semicolons, braces) to continue validation after syntax errors.")

    # Right: Clean Cropped Terminal Evidence (AST)
    img_path = screenshots_dir / "04_phase2_ast.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.41), width=Inches(6.15))


def build_slide_6_code_quality(prs: Presentation):
    """Slide 6: Code Quality & Repository Structure."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Code Quality, Modularity & Architecture")

    # 4 Architecture Cards
    cards = [
        ("1. Clean Separation of Concerns", "Every phase is isolated in a dedicated source module:\n• lexer.py (Tokenization & Lexing)\n• parser.py & ast_nodes.py (Grammar & AST)\n• symbol_table.py (Lexical Scopes)\n• semantic_analyzer.py (Type & Scope Validation)\n• tac_generator.py & interpreter.py (Backend)", BLUE_PRIMARY),
        ("2. Data-Driven Configuration", "Language syntax is decoupled from source code logic:\n• default_tokens.json (Operators, delimiters, comments)\n• keyword_mapping.json (Custom keyword dictionary)\n• config.json (Runtime behavior and path routing)\nAllows changing keywords without altering scanner logic.", GREEN_ACCENT),
        ("3. Strong Typing & Readability", "Written in Python 3.11 with dataclasses and type annotations:\n• Token dataclass with helper matchers\n• ASTVisitor abstract base class\n• Structured error classes with line/col attributes\n• Zero external package dependencies.", PURPLE_ACC),
        ("4. Transparent CLI Diagnostic Engine", "Provides interactive command-line inspection flags:\n• --tokens (Token stream table)\n• --ast (Visual ASCII AST tree)\n• --symbols (Scoped symbol table)\n• --tac (Three-Address Code IR)\n• --all (Complete pipeline trace)", ORANGE_ACC)
    ]

    for idx, (title, desc, col) in enumerate(cards):
        row = idx // 2
        col_idx = idx % 2
        left = 0.8 + col_idx * 5.95
        top = 1.6 + row * 2.7

        add_card(slide, left, top, 5.75, 2.5)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(5.75), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tbox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(5.35), Inches(2.2))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(4)

        for line in desc.splitlines():
            p_desc = tf.add_paragraph()
            p_desc.text = line
            p_desc.font.size = Pt(10.5)
            p_desc.font.color.rgb = TEXT_DARK
            p_desc.space_after = Pt(2)


def build_slide_7_testing(prs: Presentation, screenshots_dir: Path):
    """Slide 7: Representative Testing & Functional Validation."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Testing Strategy & Functional Validation")

    # Left: Representative Test Cases
    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Representative Test Cases"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = GREEN_ACCENT
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Valid Custom Script (example.custom):", "Tests Hindi keyword normalization, while-loops (jabtak), float arithmetic, and if/else branching (agar/warna).")
    add_bullet_point(tf, "• Factorial Algorithm (test_custom_hindi.custom):", "Validates loop accumulator logic, integer multiplication, and conditional verification in custom keywords.")
    add_bullet_point(tf, "• Arithmetic Precedence (test_arithmetic.edu):", "Verifies operator precedence climbing, parenthesized sub-expressions, and compound assignments (+=, -=, *=, /=).")
    add_bullet_point(tf, "• Diagnostic Error Tests (error_*.edu):", "Negative test cases validating detection of lexical errors, syntax errors, and semantic type mismatches.")

    # Right: Clean Cropped Terminal Evidence (Test Runner)
    img_path = screenshots_dir / "01_test_suite_runner.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(1.41), width=Inches(6.15))


def build_slide_8_problem_solving(prs: Presentation):
    """Slide 8: Problem Solving & Technical Innovation."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Problem Solving & Technical Innovations")

    items = [
        ("1. Multi-Character Operator Collisions", "Challenge: Operators like '=' and '==' share prefixes, leading to greedy single-character matching errors.\nSolution: Implemented length-descending operator sorting in lexer.py ensuring longest matches are prioritized deterministically.", BLUE_PRIMARY),
        ("2. Multilingual Keyword Decoupling", "Challenge: Supporting Hindi keywords (agar, jabtak) without complicating parser grammar rules.\nSolution: Lexer attaches canonical attributes (e.g. 'if' for 'agar') based on JSON mapping, allowing parser to remain grammar-pure.", GREEN_ACCENT),
        ("3. Precise Error Coordinate Tracking", "Challenge: Providing actionable feedback when syntax or semantic errors occur in nested code.\nSolution: Tracked exact 1-based line and column coordinates at token level, propagated through AST nodes to diagnostic error reporters.", PURPLE_ACC),
        ("4. Transparent Intermediate Inspection", "Challenge: Lack of visibility into internal transformations during compiler lab testing.\nSolution: Built dedicated CLI inspection options (--tokens, --ast, --symbols, --tac) allowing step-by-step verification of intermediate state.", ORANGE_ACC)
    ]

    for idx, (title, desc, col) in enumerate(items):
        row = idx // 2
        col_idx = idx % 2
        left = 0.8 + col_idx * 5.95
        top = 1.6 + row * 2.7

        add_card(slide, left, top, 5.75, 2.5)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(5.75), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tbox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(5.35), Inches(2.2))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(4)

        for line in desc.splitlines():
            p_desc = tf.add_paragraph()
            p_desc.text = line
            p_desc.font.size = Pt(10.5)
            p_desc.font.color.rgb = TEXT_DARK
            p_desc.space_after = Pt(2)


def build_slide_9_viva_qa(prs: Presentation):
    """Slide 9: Review 2 Viva & Technical Understanding."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Review 2 Viva Preparation — Key Technical Answers")

    qa_list = [
        ("Q1: What is the primary role of the Lexical Analysis stage in EduScript?",
         "A1: The Lexer converts raw character streams into categorized Token objects, strips comments and whitespace, tracks exact line/column coordinates, and normalizes custom keyword aliases into canonical tags."),
        ("Q2: How does the parser handle custom Hindi keywords without modifying its grammar?",
         "A2: The Lexer attaches canonical tags (e.g. canonical='while' for 'jabtak') based on keyword_mapping.json. The Parser checks token.canonical instead of raw lexemes, decoupling grammar rules from keyword naming."),
        ("Q3: How does the Lexer resolve ambiguities between '=' and '=='?",
         "A3: EduScript applies the maximal munch (longest match) rule by sorting the operator table descending by length. Multi-character operators like '==' are checked before single-character operators like '='."),
        ("Q4: Why is modular architecture important for compiler construction in this project?",
         "A4: Modular separation ensures each phase (Lexing, Parsing, Semantic Checking) has a single responsibility and well-defined input/output data contracts, simplifying debugging and intermediate inspection.")
    ]

    for idx, (q, a) in enumerate(qa_list):
        top = 1.6 + idx * 1.32
        add_card(slide, 0.8, top, 11.733, 1.2)

        tbox = slide.shapes.add_textbox(Inches(1.0), Inches(top + 0.08), Inches(11.333), Inches(1.0))
        tf = tbox.text_frame
        tf.word_wrap = True

        p_q = tf.paragraphs[0]
        p_q.text = q
        p_q.font.size = Pt(12)
        p_q.font.bold = True
        p_q.font.color.rgb = BLUE_PRIMARY
        p_q.space_after = Pt(2)

        p_a = tf.add_paragraph()
        p_a.text = a
        p_a.font.size = Pt(11)
        p_a.font.color.rgb = TEXT_DARK


def build_slide_10_summary(prs: Presentation, screenshots_dir: Path):
    """Slide 10: Review 2 Summary & Next Development Focus."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, "Review 2 Summary & Next Development Focus")

    # Left: Progress Summary
    add_card(slide, 0.8, 1.41, 5.2, 5.49)
    tbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.8), Inches(5.1))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Review 2 Progress Summary"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = GREEN_ACCENT
    p0.space_after = Pt(8)

    add_bullet_point(tf, "• Working Front-End Pipeline:", "Lexical analysis, keyword normalization, recursive descent parsing, and AST construction are fully functional and verified.")
    add_bullet_point(tf, "• Semantic & Scoping Base:", "Scoped symbol tables and static semantic analysis rules actively validate identifier declarations and data types.")
    add_bullet_point(tf, "• Repository & Version Control:", "https://github.com/Udbhav-Tiwari16/EduScript (Branch: main).")
    add_bullet_point(tf, "• Next Development Steps:", "Further refining semantic diagnostic edge cases, exploring intermediate representation optimizations, and finalizing documentation.")

    # Right: Clean Cropped Terminal Evidence (Git status log)
    img_path = screenshots_dir / "14_git_status_log.png"
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(6.35), Inches(2.6), width=Inches(6.15))


def main():
    base_dir = Path(__file__).parent.resolve()
    screenshots_dir = base_dir / "evidence_screenshots"
    output_ppt_path = base_dir / "EduScript_Review_2_Implementation_Progress.pptx"
    output_root_ppt_path = base_dir.parent / "EduScript_Review_2_Implementation_Progress.pptx"

    prs = create_base_presentation()

    print("Building Review 2 presentation slides...")
    build_slide_1_title(prs)
    build_slide_2_overview(prs)
    build_slide_3_progress(prs)
    build_slide_4_working_module(prs, screenshots_dir)
    build_slide_5_compiler_concepts(prs, screenshots_dir)
    build_slide_6_code_quality(prs)
    build_slide_7_testing(prs, screenshots_dir)
    build_slide_8_problem_solving(prs)
    build_slide_9_viva_qa(prs)
    build_slide_10_summary(prs, screenshots_dir)

    prs.save(str(output_ppt_path))
    prs.save(str(output_root_ppt_path))
    print(f"\nSuccessfully generated Review 2 presentation:\n- {output_ppt_path}\n- {output_root_ppt_path}")


if __name__ == "__main__":
    main()
